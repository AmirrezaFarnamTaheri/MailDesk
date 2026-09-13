from __future__ import annotations

import asyncio
import base64
import email
import email.header
import json
import os
import tempfile
import unittest
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from openpyxl import Workbook

from mailmerge_app.browser_profiles import compose_url, gmail_base_url, inbox_url
from mailmerge_app.gmail_client import (
    GMAIL_SCOPE,
    SHEETS_SCOPE,
    build_raw_message,
    new_oauth_request,
    parse_client_secret,
    spreadsheet_id_from_url,
    refresh_token,
)
from mailmerge_app.sheet_reader import list_sheets, suggest_mappings, table
from mailmerge_app.storage import Store
from mailmerge_app.template_engine import (
    batch_fingerprint,
    is_valid_email,
    message_fingerprint,
    placeholders,
    render_text,
    split_addresses,
)


class TemplateEngineTests(unittest.TestCase):
    def test_renders_defaults_conditionals_and_preserves_unknowns(self):
        rendered = render_text(
            "Hello {{ Name|there }} — {{#if Link}}Visit {{Link}}{{/if}} — {{Missing}} — {{Blank}}",
            {"Name": "Ada", "Link": "https://example.com", "Blank": ""},
            12,
        )
        self.assertEqual(rendered.text, "Hello Ada — Visit https://example.com — {{Missing}} — ")
        self.assertEqual(rendered.unresolved, ["Missing"])
        self.assertEqual(rendered.blank_values, ["Blank"])
        self.assertEqual(placeholders("{{Name}} {{ Name|friend }} {{#if Link}}{{Link}}{{/if}} {{_today}}"), ["Link", "Name", "_today"])

    def test_conditionals_are_false_for_missing_and_rtl_is_preserved(self):
        rendered = render_text("سلام {{نام|دوست}} {{#if لینک}}لینک: {{لینک}}{{/if}}", {"نام": "علیرضا"})
        self.assertEqual(rendered.text, "سلام علیرضا ")
        self.assertEqual(rendered.unresolved, [])

    def test_html_context_escapes_spreadsheet_values(self):
        rendered = render_text("<p>Hello {{Name}}</p>", {"Name": "<img src=x onerror=alert(1)>"}, escape_values=True)
        self.assertIn("&lt;img", rendered.text)
        self.assertNotIn("<img src=x", rendered.text)

    def test_email_parsing_is_conservative(self):
        self.assertTrue(is_valid_email("person@example.com"))
        self.assertFalse(is_valid_email("Person <person@example.com>"))
        self.assertFalse(is_valid_email("not-an-email"))
        self.assertEqual(split_addresses("a@example.com; b@example.com\nc@example.com"), ["a@example.com", "b@example.com", "c@example.com"])

    def test_fingerprints_include_account_mode_html_and_attachments(self):
        a = message_fingerprint("draft", "a@example.com", "x@example.com", "", "", "S", "B", "<b>B</b>", ["a"])
        b = message_fingerprint("send", "a@example.com", "x@example.com", "", "", "S", "B", "<b>B</b>", ["a"])
        c = message_fingerprint("draft", "a@example.com", "x@example.com", "", "", "S", "B", "<b>B</b>", ["b"])
        self.assertNotEqual(a, b)
        self.assertNotEqual(a, c)
        self.assertEqual(batch_fingerprint([{"to":"a@example.com","subject":"S","body":"B"}]), batch_fingerprint([{"to":"a@example.com","subject":"S","body":"B"}]))


class SheetReaderTests(unittest.TestCase):
    def test_xlsx_sheet_header_detection_and_mapping_suggestions(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "people.xlsx"
            wb = Workbook(); ws = wb.active; ws.title = "People"
            ws.append(["Mail merge source", "", ""])
            ws.append(["Full Name", "Email Address", "Status"])
            ws.append(["Ada", "ada@example.com", "Ready"])
            ws.append(["Lin", "lin@example.com", "Ready"])
            wb.save(path)
            self.assertEqual(list_sheets(path), ["People"])
            headers, rows, header_row = table(path, "People")
            self.assertEqual(header_row, 2)
            self.assertEqual(rows[0]["_row"], "3")
            suggestions = suggest_mappings(headers, rows)
            self.assertEqual(suggestions["to"], "Email Address")
            self.assertEqual(suggestions["name"], "Full Name")
            self.assertEqual(suggestions["status"], "Status")

    def test_duplicate_headers_are_disambiguated(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "duplicate.xlsx"
            wb = Workbook(); ws = wb.active; ws.title = "Sheet1"
            ws.append(["Email", "Email"]); ws.append(["one@example.com", "other@example.com"]); wb.save(path)
            headers, rows, _ = table(path, "Sheet1", 1)
            self.assertEqual(headers, ["Email", "Email (2)"])
            self.assertEqual(rows[0]["Email (2)"], "other@example.com")

    def test_large_sheet_is_read_without_truncating_reasonable_batches(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "large.xlsx"
            wb = Workbook(write_only=True); ws = wb.create_sheet("People"); ws.append(["Name", "Email"])
            for i in range(5000): ws.append([f"Person {i}", f"p{i}@example.com"])
            wb.save(path)
            _, rows, _ = table(path, "People", 1)
            self.assertEqual(len(rows), 5000)


class GmailClientTests(unittest.TestCase):
    def test_desktop_client_secret_only(self):
        good = json.dumps({"installed": {"client_id": "id", "client_secret": "secret"}}).encode()
        self.assertEqual(parse_client_secret(good), {"client_id": "id", "client_secret": "secret"})
        bad = json.dumps({"web": {"client_id": "id", "client_secret": "secret"}}).encode()
        with self.assertRaises(ValueError): parse_client_secret(bad)

    def test_oauth_can_request_gmail_and_sheets_scopes_with_pkce(self):
        url, session = new_oauth_request({"client_id": "id", "client_secret": "secret"}, "http://127.0.0.1:8765/oauth/google/callback", scopes=[GMAIL_SCOPE, SHEETS_SCOPE])
        query = parse_qs(urlparse(url).query)
        self.assertIn(GMAIL_SCOPE, query["scope"][0])
        self.assertIn(SHEETS_SCOPE, query["scope"][0])
        self.assertEqual(query["code_challenge_method"], ["S256"])
        self.assertEqual(query["state"], [session["state"]])

    def test_mime_message_contains_html_and_attachment(self):
        raw = build_raw_message("to@example.com", "Résumé ✓", "Hello", "cc@example.com", "", "<b>Hello</b>", [("note.txt", b"hello", "text/plain")])
        padded = raw + "=" * (-len(raw) % 4)
        parsed = email.message_from_bytes(base64.urlsafe_b64decode(padded.encode("ascii")))
        self.assertEqual(parsed["To"], "to@example.com")
        self.assertEqual(str(email.header.make_header(email.header.decode_header(parsed["Subject"]))), "Résumé ✓")
        self.assertTrue(parsed.is_multipart())
        self.assertIn("note.txt", [part.get_filename() for part in parsed.walk() if part.get_filename()])

    def test_google_sheet_id_parsing(self):
        sid = "1AbCdEf_123-xyz"
        self.assertEqual(spreadsheet_id_from_url(f"https://docs.google.com/spreadsheets/d/{sid}/edit#gid=0"), sid)
        self.assertEqual(spreadsheet_id_from_url(sid), sid)
        with self.assertRaises(ValueError): spreadsheet_id_from_url("not a valid id / url")

    def test_expired_oauth_without_refresh_token_requires_reconnect(self):
        token = {"access_token": "old", "expires_at": "2000-01-01T00:00:00+00:00"}
        with self.assertRaises(RuntimeError):
            asyncio.run(refresh_token(token, {"client_id": "id", "client_secret": "secret"}))


class BrowserComposeTests(unittest.TestCase):
    def test_compose_url_targets_exact_gmail_slot(self):
        url = compose_url("to@example.com", "Hello & hi", "Line 1\nLine 2", gmail_slot=1)
        parsed = urlparse(url)
        query = parse_qs(parsed.query)
        self.assertEqual(parsed.path, "/mail/u/1/")
        self.assertEqual(query["to"], ["to@example.com"])
        self.assertEqual(query["su"], ["Hello & hi"])
        self.assertEqual(query["view"], ["cm"])
        self.assertEqual(gmail_base_url(0), "https://mail.google.com/mail/u/0/")
        self.assertEqual(inbox_url(1), "https://mail.google.com/mail/u/1/#inbox")


class StoreTests(unittest.TestCase):
    def test_templates_browser_routes_queue_history_and_restart_pause(self):
        with tempfile.TemporaryDirectory() as td:
            old = os.environ.get("MAILMERGE_DATA_DIR"); os.environ["MAILMERGE_DATA_DIR"] = td
            try:
                path = Path(td) / "test.db"; store = Store(path)
                saved = store.save_template({"id":"test","name":"Test","subject":"Hi {{Name}}","body":"Body","body_html":"<b>Body</b>","signature_html":"","cc_template":"","bcc_template":"","attachment_ids":[],"tags":"test","default_account":"","default_browser_sender_id":""})
                self.assertEqual(saved["body_html"], "<b>Body</b>")
                route = store.save_browser_sender({"id":"route","label":"Work","browser_id":"chrome","profile_dir":"Default","gmail_slot":1,"expected_email":"work@example.com"})
                self.assertEqual(route["gmail_slot"], 1)
                campaign = store.create_campaign({"name":"Dry","mode":"dry_run","batch_id":"abc123456789","status":"Queued"}, [{"row_number":"2","to":"to@example.com","subject":"Hi","body":"Body","fingerprint":"fp","attachments":[]}])
                store.set_campaign_status(campaign["id"], "Running")
                reopened = Store(path)
                self.assertEqual(reopened.get_campaign(campaign["id"])["status"], "Paused")
                reopened.log_operation({"campaign_id":campaign["id"],"mode":"dry_run","account":"","row_number":"2","recipient":"to@example.com","subject":"Hi","fingerprint":"fp","remote_id":"DryRunOnly","result":"Success","error":""})
                self.assertTrue(reopened.fingerprint_succeeded("fp"))
            finally:
                if old is None: os.environ.pop("MAILMERGE_DATA_DIR", None)
                else: os.environ["MAILMERGE_DATA_DIR"] = old


if __name__ == "__main__":
    unittest.main()
