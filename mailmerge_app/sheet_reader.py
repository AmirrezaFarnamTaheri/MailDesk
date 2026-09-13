from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm", ".csv"}
MAX_ROWS = 100_000
MAX_CSV_FIELD_BYTES = 4 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 300 * 1024 * 1024
MAX_XLSX_ENTRIES = 20_000


def _display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _validate_xlsx_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_XLSX_ENTRIES:
                raise ValueError(f"Spreadsheet archive contains too many files ({len(infos):,}).")
            uncompressed = sum(info.file_size for info in infos)
            if uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ValueError(
                    f"Spreadsheet expands to {uncompressed / 1024 / 1024:.1f} MB; the safety limit is "
                    f"{MAX_XLSX_UNCOMPRESSED_BYTES // 1024 // 1024} MB."
                )
            if any(info.flag_bits & 0x1 for info in infos):
                raise ValueError("Encrypted spreadsheet archives are not supported.")
    except zipfile.BadZipFile as exc:
        raise ValueError("The spreadsheet is not a valid XLSX/XLSM archive.") from exc


def list_sheets(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return [path.stem]
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported spreadsheet type: {suffix}")
    _validate_xlsx_archive(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _csv_rows(path: Path) -> list[list[str]]:
    raw = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "utf-8", "cp1256", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode("utf-8", errors="replace")

    # csv.reader must receive the original newline stream. splitlines() corrupts
    # valid quoted fields containing embedded newlines by turning one logical row
    # into several physical rows.
    previous_limit = csv.field_size_limit()
    csv.field_size_limit(max(previous_limit, MAX_CSV_FIELD_BYTES))
    rows: list[list[str]] = []
    try:
        reader = csv.reader(io.StringIO(text, newline=""))
        for index, row in enumerate(reader):
            if index >= MAX_ROWS + 100:
                raise ValueError(f"Spreadsheet exceeds the {MAX_ROWS:,}-row safety limit.")
            rows.append([cell.strip() for cell in row])
    except csv.Error as exc:
        raise ValueError(f"Could not parse CSV: {exc}") from exc
    finally:
        csv.field_size_limit(previous_limit)
    return rows


def _xlsx_rows(path: Path, sheet: str) -> list[list[str]]:
    _validate_xlsx_archive(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"Worksheet not found: {sheet}")
        ws = wb[sheet]
        rows: list[list[str]] = []
        for index, row in enumerate(ws.iter_rows(values_only=True)):
            if index >= MAX_ROWS + 100:
                raise ValueError(f"Spreadsheet exceeds the {MAX_ROWS:,}-row safety limit.")
            rows.append([_display_value(cell) for cell in row])
        return rows
    finally:
        wb.close()


def raw_rows(path: Path, sheet: str) -> list[list[str]]:
    return _csv_rows(path) if path.suffix.lower() == ".csv" else _xlsx_rows(path, sheet)


def detect_header_row(rows: list[list[str]], scan_limit: int = 30) -> int:
    best_row = 1
    best_score = -1
    for index, row in enumerate(rows[:scan_limit], start=1):
        nonempty = [cell.strip() for cell in row if str(cell).strip()]
        unique = len(set(nonempty))
        score = unique * 3 - max(0, len(nonempty) - unique) * 2
        if len(nonempty) >= 2 and score > best_score:
            best_row = index
            best_score = score
    return best_row


def table(path: Path, sheet: str, header_row: int | None = None) -> tuple[list[str], list[dict[str, str]], int]:
    rows = raw_rows(path, sheet)
    if not rows:
        return [], [], 1
    selected_header = header_row or detect_header_row(rows)
    if selected_header < 1 or selected_header > len(rows):
        raise ValueError("Header row is outside the spreadsheet.")
    if len(rows) - selected_header > MAX_ROWS:
        raise ValueError(f"Spreadsheet exceeds the {MAX_ROWS:,}-data-row safety limit.")

    header_cells = rows[selected_header - 1]
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, cell in enumerate(header_cells, start=1):
        base = cell.strip() or f"Column {index}"
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base} ({counts[base]})")

    records: list[dict[str, str]] = []
    for excel_row, row in enumerate(rows[selected_header:], start=selected_header + 1):
        padded = list(row) + [""] * max(0, len(headers) - len(row))
        record = {header: (padded[i] if i < len(padded) else "") for i, header in enumerate(headers)}
        record["_row"] = str(excel_row)
        if any(value.strip() for key, value in record.items() if key != "_row"):
            records.append(record)
    return headers, records, selected_header


def column_profiles(headers: list[str], rows: list[dict[str, str]], sample_size: int = 200) -> list[dict[str, Any]]:
    sample = rows[:sample_size]
    profiles: list[dict[str, Any]] = []
    email_re = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
    for header in headers:
        values = [row.get(header, "").strip() for row in sample]
        nonempty = [value for value in values if value]
        emails = [value for value in nonempty if email_re.match(value)]
        unique = len(set(nonempty))
        profiles.append(
            {
                "name": header,
                "nonempty": len(nonempty),
                "unique": unique,
                "email_ratio": round(len(emails) / len(nonempty), 3) if nonempty else 0,
                "examples": nonempty[:3],
            }
        )
    return profiles


def suggest_mappings(headers: list[str], rows: list[dict[str, str]]) -> dict[str, str]:
    profiles = column_profiles(headers, rows)
    normalized = {header: re.sub(r"[^a-z0-9]+", "", header.casefold()) for header in headers}

    def by_name(words: tuple[str, ...]) -> str:
        for header, key in normalized.items():
            if any(word in key for word in words):
                return header
        return ""

    email = max(profiles, key=lambda p: p["email_ratio"], default={"name": "", "email_ratio": 0})
    return {
        "to": email["name"] if email.get("email_ratio", 0) >= 0.5 else by_name(("email", "mail", "ایمیل")),
        "name": by_name(("firstname", "fullname", "name", "student", "نام")),
        "cc": by_name(("cc", "copy")),
        "bcc": by_name(("bcc",)),
        "attachment": by_name(("attachment", "file", "document", "پیوست", "فایل")),
        "status": by_name(("status", "state", "وضعیت")),
    }
